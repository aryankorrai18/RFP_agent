from __future__ import annotations

import io

import docx
import openpyxl
import pytest

from backend.parser import MIN_PDF_TEXT_CHARS, ParseError, parse_document
from tests.pdf_helpers import blank_pdf, text_pdf

LIMIT = 400_000


def test_text_file_strips_bom():
    doc = parse_document("rfp.txt", "﻿1. Do you support SSO?".encode(), LIMIT)
    assert doc.kind == "text"
    assert doc.text == "1. Do you support SSO?"
    assert doc.pdf_bytes is None


def test_markdown_is_read_as_text():
    assert parse_document("rfp.md", b"# Security\nDo you encrypt data?", LIMIT).kind == "text"


def test_docx_keeps_headings_paragraphs_and_tables_in_order():
    document = docx.Document()
    document.add_heading("3. Security", level=1)
    document.add_paragraph("3.1 Do you support SSO?")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Requirement", "Compliant"
    table.cell(1, 0).text, table.cell(1, 1).text = "RBAC must be supported", ""
    document.add_paragraph("3.2 Describe encryption.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = parse_document("rfp.docx", buffer.getvalue(), LIMIT).text
    assert text.splitlines() == [
        "## 3. Security",
        "3.1 Do you support SSO?",
        "Requirement | Compliant",
        "RBAC must be supported",
        "3.2 Describe encryption.",
    ]


def test_xlsx_lists_every_sheet_and_non_empty_row():
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Security"
    first.append(["ID", "Question"])
    first.append(["SEC-01", "Do you support SAML?"])
    first.append([None, None])
    second = workbook.create_sheet("Privacy")
    second.append(["PRV-01", "Do you sign a DPA?"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = parse_document("q.xlsx", buffer.getvalue(), LIMIT).text
    assert text.splitlines() == [
        "[Sheet: Security]",
        "Row 1: ID | Question",
        "Row 2: SEC-01 | Do you support SAML?",
        "[Sheet: Privacy]",
        "Row 1: PRV-01 | Do you sign a DPA?",
    ]


def test_pdf_with_text_is_extracted_with_page_markers():
    lines = [f"3.{i} Describe your approach to requirement number {i} in detail." for i in range(1, 8)]
    doc = parse_document("rfp.pdf", text_pdf(lines), LIMIT)
    assert doc.pdf_bytes is None
    assert doc.text.startswith("[Page 1]")
    assert "3.1 Describe your approach" in doc.text


def test_pdf_without_text_falls_back_to_sending_the_pdf():
    data = blank_pdf()
    doc = parse_document("scan.pdf", data, LIMIT)
    assert doc.pdf_bytes == data
    assert len(doc.text) < MIN_PDF_TEXT_CHARS


@pytest.mark.parametrize("name", ["rfp.doc", "rfp.pptx", "rfp", "rfp.csv"])
def test_unsupported_extensions_are_rejected(name):
    with pytest.raises(ParseError) as error:
        parse_document(name, b"data", LIMIT)
    assert error.value.code == "unsupported_file_type"


def test_corrupt_docx_is_unreadable():
    with pytest.raises(ParseError) as error:
        parse_document("rfp.docx", b"not a zip file", LIMIT)
    assert error.value.code == "unreadable_file"


def test_empty_text_file_is_rejected():
    with pytest.raises(ParseError) as error:
        parse_document("rfp.txt", b"   \n  ", LIMIT)
    assert error.value.code == "empty_document"


def test_empty_workbook_is_rejected():
    buffer = io.BytesIO()
    openpyxl.Workbook().save(buffer)
    with pytest.raises(ParseError) as error:
        parse_document("q.xlsx", buffer.getvalue(), LIMIT)
    assert error.value.code == "empty_document"


def test_document_over_the_character_limit_is_rejected():
    with pytest.raises(ParseError) as error:
        parse_document("rfp.txt", b"x" * 101, 100)
    assert error.value.code == "document_too_long"
