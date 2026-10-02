"""Export approved answers (design §9). Blocked until every requirement is final (PRD NFR-02).

- DOCX: a heading per section, then each requirement's reference, question and final answer.
- XLSX: the buyer's original workbook with answers written into it. Each requirement's row is
  found by matching its question text; the answer goes into a response/answer/comment column
  (or a new "Vendor response" column). Anything that can't be placed is listed on an extra
  "Unplaced answers" sheet instead of being dropped silently.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

import docx
import openpyxl

ANSWER_HEADER = re.compile(r"response|answer|comment|vendor|reply", re.I)


@dataclass(frozen=True)
class FinalAnswer:
    code: str
    section: str | None
    reference: str | None
    question: str
    answer: str


def export_docx(title: str, subtitle: str | None, answers: list[FinalAnswer]) -> bytes:
    document = docx.Document()
    document.add_heading(title, level=0)
    if subtitle:
        document.add_paragraph(subtitle)
    current_section = object()
    for item in answers:
        if item.section != current_section:
            current_section = item.section
            if item.section:
                document.add_heading(item.section, level=1)
        question = document.add_paragraph()
        question.add_run(f"{item.reference or item.code} ").bold = True
        question.add_run(item.question).bold = True
        document.add_paragraph(item.answer)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def export_xlsx(original: Path, answers: list[FinalAnswer]) -> tuple[bytes, list[str]]:
    """Returns the filled workbook and the codes of answers that couldn't be placed."""
    workbook = openpyxl.load_workbook(original)
    unplaced: list[FinalAnswer] = []
    answer_columns: dict[str, int] = {}

    for item in answers:
        target = _find_question_cell(workbook, item.question)
        if target is None:
            unplaced.append(item)
            continue
        sheet, row, question_column = target
        column = answer_columns.get(sheet.title) or _answer_column(sheet, question_column)
        answer_columns[sheet.title] = column
        sheet.cell(row=row, column=column, value=item.answer)

    if unplaced:
        extra = workbook.create_sheet("Unplaced answers")
        extra.append(["Reference", "Question", "Answer"])
        for item in unplaced:
            extra.append([item.reference or item.code, item.question, item.answer])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), [item.code for item in unplaced]


def _find_question_cell(workbook, question: str):  # noqa: ANN001, ANN202
    wanted = _norm(question)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and _norm(cell.value) == wanted:
                    return sheet, cell.row, cell.column
    # Fall back to containment (the extractor may have trimmed a prefix such as a number).
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and wanted and wanted in _norm(cell.value):
                    return sheet, cell.row, cell.column
    return None


def _answer_column(sheet, question_column: int) -> int:  # noqa: ANN001
    header_rows = range(1, min(sheet.max_row, 5) + 1)
    for row in header_rows:
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row=row, column=column).value
            if isinstance(value, str) and ANSWER_HEADER.search(value) and column != question_column:
                return column
    column = sheet.max_column + 1
    sheet.cell(row=1, column=column, value="Vendor response")
    return column
