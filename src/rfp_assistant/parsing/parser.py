"""Turn an uploaded file into plain text for requirement extraction. No AI here."""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import PurePath

import docx
import openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".txt": "text", ".md": "text"}

# Below this much extracted text, a PDF is treated as scanned and sent to the model as a PDF.
MIN_PDF_TEXT_CHARS = 200


class ParseError(Exception):
    """A file that can't be turned into usable text. `code` becomes the API error code."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    kind: str
    text: str
    # Set only for PDFs with too little extractable text; extraction then sends the PDF itself.
    pdf_bytes: bytes | None = None

    @property
    def char_count(self) -> int:
        return len(self.text)


def parse_document(filename: str, data: bytes, max_chars: int) -> ParsedDocument:
    suffix = PurePath(filename).suffix.lower()
    kind = SUPPORTED_EXTENSIONS.get(suffix)
    if kind is None:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ParseError("unsupported_file_type", f"'{suffix or filename}' is not supported. Use one of: {supported}.")

    try:
        if kind == "pdf":
            text = _pdf_text(data)
        elif kind == "docx":
            text = _docx_text(data)
        elif kind == "xlsx":
            text = _xlsx_text(data)
        else:
            text = data.decode("utf-8", errors="replace").lstrip("﻿")
    except ParseError:
        raise
    except Exception as exc:  # the parsing libraries raise many different exception types
        raise ParseError("unreadable_file", f"Could not read {filename} as {kind.upper()}: {exc}") from exc

    text = text.strip()
    if len(text) > max_chars:
        raise ParseError(
            "document_too_long",
            f"{filename} contains {len(text):,} characters of text; the limit is {max_chars:,}.",
        )

    if kind == "pdf" and len(text) < MIN_PDF_TEXT_CHARS:
        return ParsedDocument(filename=filename, kind=kind, text=text, pdf_bytes=data)
    if not text:
        raise ParseError("empty_document", f"No text was found in {filename}.")
    return ParsedDocument(filename=filename, kind=kind, text=text)


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ParseError("unreadable_file", "The PDF is password-protected.")
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(f"[Page {number}]\n{page_text}")
    return "\n\n".join(pages)


def _docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    # iter_inner_content keeps paragraphs and tables in document order, which matters
    # because RFPs often put questions in tables under a section heading.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = block.style.name if block.style is not None else ""
            is_heading = style.lower().startswith("heading") or style == "Title"
            lines.append(f"## {text}" if is_heading else text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    # Merged cells are repeated by python-docx; keep one copy.
                    if cell_text and (not cells or cells[-1] != cell_text):
                        cells.append(cell_text)
                if cells:
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def _xlsx_text(data: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    has_rows = False
    try:
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    has_rows = True
                    lines.append(f"Row {row_number}: " + " | ".join(values))
    finally:
        workbook.close()
    return "\n".join(lines) if has_rows else ""
