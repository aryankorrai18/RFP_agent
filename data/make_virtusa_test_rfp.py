"""Build a fictional test RFP .docx for the Virtusa workspace. Its questions deliberately overlap
the topics of both past proposals (one won, one lost), so drafting has to choose between the
winning, specific answer and the losing, vague one — the core outcome-aware-ranking test."""

import docx

SECTIONS = [
    ("2. Company", [
        "Provide an overview of your company, including headcount and locations.",
    ]),
    ("3. Identity and access", [
        "Which single sign-on protocols and identity providers do you support?",
        "How are user accounts created and removed automatically when staff join or leave?",
        "Is multi-factor authentication enforced for administrative users?",
    ]),
    ("4. Security assurance", [
        "List your current security certifications and attestations.",
        "How often is your platform tested, and by whom?",
        "How long are audit logs retained, and can they be exported?",
    ]),
    ("5. Hosting and service", [
        "Where is client data stored?",
        "What availability do you commit to, and what remedies apply if it is missed?",
        "Describe your support model, including response times for critical incidents.",
    ]),
    ("6. Delivery", [
        "Describe your implementation approach and a typical timeline.",
    ]),
]


def build(path):
    document = docx.Document()
    document.add_heading("Meridian Health Group: Request for Proposal (fictional test document)", level=0)
    document.add_paragraph("RFP reference TEST-RFP-2026-002 | Fictional test document, for trying the drafting flow only")
    document.add_heading("1. Instructions", level=1)
    document.add_paragraph("Answer every question below.")
    for heading, questions in SECTIONS:
        document.add_heading(heading, level=1)
        for question in questions:
            document.add_paragraph(question)
    document.save(path)


if __name__ == "__main__":
    build("data/virtusa_test_rfp.docx")
    print("wrote data/virtusa_test_rfp.docx")
