"""Build a second fictional test past-proposal .docx for the Virtusa workspace: a LOST proposal
with vaguer answers on the same topics as the won one, so outcome-aware ranking has something to
prefer between. Not a real Virtusa document."""

import docx

QA = [
    ("Company overview", [
        ("Tell us about your company, including size and locations.",
         "Virtusa is an IT services company with a global presence and a large team."),
    ]),
    ("Identity and access", [
        ("Which single sign-on protocols and identity providers do you support?",
         "We support industry-standard single sign-on and can work with most identity providers."),
        ("How are user accounts created and removed automatically when staff join or leave?",
         "User accounts are managed by our support team on request. Automated provisioning is on our roadmap."),
        ("Is multi-factor authentication enforced for administrative users?",
         "Multi-factor authentication is available as an option for administrative accounts."),
    ]),
    ("Security assurance", [
        ("List your current security certifications and attestations.",
         "We follow industry-standard security practices and are working toward formal certification."),
        ("How often is your platform tested, and by whom?",
         "We perform regular security testing of our platform and take security very seriously."),
        ("How long are audit logs retained, and can they be exported?",
         "Audit logs are kept for a reasonable period and can be provided on request."),
    ]),
    ("Hosting and service", [
        ("Where is client data stored?",
         "Client data is stored in secure cloud infrastructure."),
        ("What availability do you commit to, and what remedies apply if it is missed?",
         "We aim for high availability. Specific SLA terms are discussed during contracting."),
        ("Describe your support model, including response times for critical incidents.",
         "Support is available during business hours, with escalation for urgent issues."),
    ]),
    ("Delivery", [
        ("Describe your implementation approach and a typical timeline.",
         "Every engagement is different, so we agree the plan and timeline during discovery with the client."),
    ]),
]


def build(path):
    document = docx.Document()
    document.add_heading("Virtusa: Response to Northbridge Insurance RFP (fictional test document)", level=0)
    document.add_paragraph("RFP reference TEST-2025-014 | Submitted 2025-11-03 | Prepared by Virtusa "
                           "Proposals Team | For testing the import flow only")
    document.add_paragraph("Virtusa thanks you for the opportunity to respond. Our answers to each "
                           "question follow, in the order of your request.")
    number = 1
    for heading, items in QA:
        document.add_heading(heading, level=1)
        for question, answer in items:
            document.add_paragraph(f"Q{number}. {question}")
            document.add_paragraph(f"Response: {answer}")
            number += 1
    document.save(path)


if __name__ == "__main__":
    build("data/virtusa_test_past_proposal_lost.docx")
    print("wrote data/virtusa_test_past_proposal_lost.docx")
