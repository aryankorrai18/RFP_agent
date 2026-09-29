"""Build a fictional test past-proposal .docx for the Virtusa workspace, so you can try the
Library -> import flow. Not a real Virtusa document."""

import docx

QA = [
    ("Company overview", [
        ("Provide an overview of your company, including headcount and locations.",
         "Virtusa is a global digital engineering and IT services provider, founded in 1996 and "
         "headquartered in Southborough, Massachusetts. We employ approximately 30,000 people "
         "across delivery centers in India, Sri Lanka, and other locations worldwide."),
    ]),
    ("Identity and access", [
        ("Which single sign-on protocols and identity providers do you support?",
         "We support SAML 2.0 and OpenID Connect, tested with Okta, Microsoft Entra ID, and Google Workspace."),
        ("How are user accounts created and removed automatically when staff join or leave?",
         "We support automatic user provisioning and de-provisioning via SCIM 2.0, integrated with "
         "the client's identity provider. De-provisioned accounts lose access within minutes."),
        ("Is multi-factor authentication enforced for administrative users?",
         "Yes. Multi-factor authentication is mandatory for all administrative accounts."),
    ]),
    ("Security assurance", [
        ("List your current security certifications and attestations.",
         "We hold ISO/IEC 27001 certification and complete an annual SOC 2 Type II audit."),
        ("How often is your platform tested, and by whom?",
         "An independent third-party firm performs a full penetration test at least once a year, "
         "with quarterly automated vulnerability scans in between. An executive summary is "
         "available to clients under NDA."),
        ("How long are audit logs retained, and can they be exported?",
         "Audit logs of user and administrator activity are retained for 12 months and can be "
         "exported on request via the client portal."),
    ]),
    ("Hosting and service", [
        ("Where is client data stored?",
         "Client data is hosted on AWS, in the region the client selects at onboarding (US or EU)."),
        ("What availability do you commit to, and what remedies apply if it is missed?",
         "Our standard SLA commits to 99.9% monthly uptime for production services, with service "
         "credits applied if that commitment is missed."),
        ("Describe your support model, including response times for critical incidents.",
         "Severity-1 incidents receive 24/7 coverage with an initial response within 30 minutes. "
         "Other severities are handled during business hours via a support portal and email."),
    ]),
    ("Delivery", [
        ("Describe your implementation approach and a typical timeline.",
         "Implementations follow a phased approach: discovery (weeks 1-2), configuration and "
         "integration (weeks 3-5), user acceptance testing (week 6), and go-live with hypercare "
         "support. A typical engagement takes 6 to 8 weeks depending on scope."),
    ]),
]


def build(path):
    document = docx.Document()
    document.add_heading("Virtusa: Response to Example Bank RFP (fictional test document)", level=0)
    document.add_paragraph("RFP reference TEST-2026-001 | Submitted 2026-06-15 | Prepared by Virtusa "
                           "Proposals Team | For testing the import flow only")
    document.add_paragraph("Virtusa thanks you for the opportunity to respond. Our answers to each "
                           "question follow, in the order of your request.")
    number = 1
    for heading, items in QA:
        document.add_heading(heading, level=1)
        for question, answer in items:
            document.add_paragraph(f"Q{number}. {question}")
            document.add_paragraph(f"Response: {answer}")
            number += 1
    document.save(path)


if __name__ == "__main__":
    build("data/virtusa_test_past_proposal.docx")
    print("wrote data/virtusa_test_past_proposal.docx")
