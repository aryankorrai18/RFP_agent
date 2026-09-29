"""Generate the sample documents used by the V0 acceptance test.

Everything here is fictional: the buyer (Harborview Credit Union) and the vendor
(Larkspur Data, see data/fact_sheet.json).

Some questions are deliberately NOT covered by the fact sheet (SCIM provisioning, MFA,
customer references, a core-banking integration, pricing). A correct V0 marks those
"Needs SME" instead of inventing an answer.

Run:  python samples/make_samples.py
"""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent

RFP_SECTIONS: list[tuple[str, list[tuple[str, str]]]] = [
    (
        "2. Company Information",
        [
            ("2.1", "Provide an overview of your company, including year founded, number of employees and office locations. (Maximum 150 words)"),
            ("2.2", "Provide three customer references from financial-services organisations, with contact details."),
        ],
    ),
    (
        "3. Security",
        [
            ("3.1", "MANDATORY: Describe how your solution supports single sign-on with our identity provider (Microsoft Entra ID) via SAML 2.0. (Maximum 100 words)"),
            ("3.2", "MANDATORY: Does your solution support automatic user provisioning and de-provisioning via SCIM 2.0?"),
            ("3.3", "Describe how customer data is encrypted at rest and in transit."),
            ("3.4", "Provide your most recent SOC 2 Type II report or describe your current certifications."),
            ("3.5", "How often is your platform penetration tested, and by whom?"),
        ],
    ),
    (
        "4. Data and Hosting",
        [
            ("4.1", "MANDATORY: Where will our data be stored? Confirm whether data can be kept within the European Union."),
            ("4.2", "Describe your backup and disaster-recovery arrangements, including RPO and RTO."),
            ("4.3", "What happens to our data when the contract ends?"),
        ],
    ),
    (
        "5. Service and Support",
        [
            ("5.1", "State your uptime commitment and what remedies apply if it is not met."),
            ("5.2", "Describe your support hours and channels, including for critical incidents."),
            ("5.3", "Describe your implementation approach and a typical timeline. (Maximum 120 words)"),
        ],
    ),
    (
        "6. Integration",
        [
            ("6.1", "Describe the APIs and prebuilt connectors your platform provides."),
            ("6.2", "Can your platform ingest real-time transaction events from our core banking system (Jack Henry Symitar)?"),
        ],
    ),
    (
        "7. Pricing",
        [
            ("7.1", "Provide a three-year pricing proposal, including implementation fees."),
        ],
    ),
]

COMPLIANCE_TABLE = [
    ("Requirement", "Compliant (Y/N)", "Comments"),
    ("8.1 The solution must provide role-based access control.", "", ""),
    ("8.2 Audit logs must be retained for at least 12 months.", "", ""),
    ("8.3 Multi-factor authentication must be enforced for all administrative access.", "", ""),
    ("8.4 The vendor must notify us of any confirmed security breach within 72 hours.", "", ""),
]

QUESTIONNAIRE = [
    ("SEC-01", "Access control", "Do you support SAML 2.0 single sign-on?"),
    ("SEC-02", "Access control", "Do you support role-based access control with custom roles?"),
    ("SEC-03", "Access control", "Is multi-factor authentication enforced for all user accounts?"),
    ("SEC-04", "Encryption", "Which encryption standard protects data at rest?"),
    ("SEC-05", "Encryption", "Which TLS versions are supported for data in transit?"),
    ("SEC-06", "Assurance", "Do you hold an ISO/IEC 27001 certification? If so, when was it issued?"),
    ("SEC-07", "Assurance", "How often are independent penetration tests performed?"),
    ("SEC-08", "Resilience", "What are your RPO and RTO?"),
    ("SEC-09", "Privacy", "Do you offer a GDPR data processing agreement?"),
    ("SEC-10", "Privacy", "Do you use subprocessors located outside the EU?"),
]


# --- V1: past proposals for the answer library (fictional) -------------------------------------
# Each covers questions the sample RFP asks, in different words, so retrieval is tested by meaning.
# Some answers add knowledge the fact sheet lacks (SCIM, admin MFA, the pen-test firm,
# implementation phases, service-credit tiers), so V1 can ground answers V0 had to send to an SME.
# Two answers are deliberately outdated and conflict with the current fact sheet (employee count,
# TLS version): the fact sheet must win.

PAST_PROPOSALS = {
    "past_northwind_bank_2025.docx": (
        "Larkspur Data: Response to Northwind Bank RFP NB-2025-07 (fictional)",
        [
            ("2. About Larkspur", [
                ("2.1", "Tell us about your organisation.",
                 "Larkspur Data was founded in 2019 and provides a cloud customer-data platform built for regulated "
                 "industries such as banking and insurance. Our team of around 200 people works from Denver and Dublin."),
            ]),
            ("3. Identity and Security", [
                ("3.1", "Which single sign-on standards and identity providers do you support?",
                 "We support single sign-on using SAML 2.0 and OpenID Connect. We have tested SSO with Microsoft Entra "
                 "ID, Okta and Google Workspace, and customers can require SSO for all users."),
                ("3.2", "Describe how user accounts are created and removed automatically.",
                 "Larkspur supports automatic user provisioning and de-provisioning through SCIM 2.0 with Okta and "
                 "Microsoft Entra ID. Deactivating a user in the identity provider removes their Larkspur access within "
                 "minutes."),
                ("3.3", "How do you protect customer data cryptographically?",
                 "Customer data is encrypted at rest with AES-256 using keys managed in AWS KMS, and in transit with "
                 "TLS 1.2 or higher."),
            ]),
            ("5. Delivery", [
                ("5.3", "Outline your onboarding methodology.",
                 "Implementation follows four phases: discovery (weeks 1-2), configuration and data connection "
                 "(weeks 3-5), user acceptance testing (week 6), and go-live with hypercare (weeks 7-8). A dedicated "
                 "implementation manager leads the project throughout."),
            ]),
            ("6. Integration", [
                ("6.1", "What integration options does the platform offer?",
                 "We provide a documented REST API and prebuilt connectors for Snowflake, Google BigQuery, Amazon S3 "
                 "and Salesforce."),
            ]),
        ],
    ),
    "past_meridian_health_2026.docx": (
        "Larkspur Data: Proposal to Meridian Health Partners (fictional)",
        [
            ("Security", [
                ("S-4", "Who performs penetration testing of your service, and how often?",
                 "Our platform is penetration tested at least annually by Ironbridge Security, an independent "
                 "CREST-accredited firm. We share the executive summary of the latest report under NDA."),
                ("S-5", "Is multi-factor authentication required for administrators?",
                 "Multi-factor authentication is enforced for all administrative accounts. Administrators sign in through "
                 "SSO with MFA required by the identity provider, or with Larkspur's built-in TOTP MFA."),
                ("S-8", "How long are audit logs kept?",
                 "Audit logs of user and administrator activity are retained for 12 months and can be exported through "
                 "our API."),
            ]),
            ("Data", [
                ("D-1", "Can all patient-related data remain within the European Union?",
                 "Customer data can be hosted entirely within the EU in our AWS Frankfurt region (eu-central-1). Data, "
                 "backups and logs remain in the EU."),
                ("D-3", "Explain your data handling at contract termination.",
                 "When the contract ends we delete all customer data within 30 days, after offering a full export."),
            ]),
        ],
    ),
    "past_cobalt_insurance_2025.docx": (
        "Larkspur Data: Response to Cobalt Insurance security and service questionnaire (fictional)",
        [
            ("Service levels", [
                ("SL-1", "What availability do you guarantee, and what happens if you miss it?",
                 "We commit to 99.9% monthly availability. If we miss it, customers receive service credits of 10% of "
                 "the monthly fee below 99.9% and 25% below 99.0%."),
                ("SL-2", "When and how can we reach support?",
                 "Severity-1 incidents are handled 24/7. Other requests are handled Monday to Friday, 8am-6pm in the "
                 "customer's region, through our support portal or by email."),
            ]),
            ("Security", [
                ("SEC-3", "Is data encrypted?",
                 "All data is encrypted in transit using TLS 1.1 or higher and at rest with AES-256."),
                ("SEC-6", "Describe backup and recovery.",
                 "Backups are taken daily, encrypted, kept for 35 days and replicated to a second region in the same "
                 "geography. Our RPO is 1 hour and our RTO is 4 hours."),
            ]),
        ],
    ),
}


def make_past_proposal(path: Path, title: str, sections) -> None:  # noqa: ANN001
    document = docx.Document()
    document.add_heading(title, level=0)
    for heading, items in sections:
        document.add_heading(heading, level=1)
        for number, question, answer in items:
            document.add_paragraph(f"{number} {question}")
            document.add_paragraph(f"Response: {answer}")
    document.save(path)


def make_rfp(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Harborview Credit Union: Request for Proposal", level=0)
    document.add_paragraph("Customer Data Platform, RFP reference HCU-2026-014 (fictional)")

    document.add_heading("1. Instructions to Vendors", level=1)
    document.add_paragraph("Proposals must be submitted by 30 November 2026, 5pm Eastern, to procurement@harborview.example.")
    document.add_paragraph("Answer every question in sections 2 to 8. Where a word limit is stated, answers over the limit may not be scored.")

    for heading, questions in RFP_SECTIONS:
        document.add_heading(heading, level=1)
        for number, text in questions:
            document.add_paragraph(f"{number} {text}")

    document.add_heading("8. Compliance Matrix", level=1)
    document.add_paragraph("Indicate compliance with each requirement below and add comments where needed.")
    table = document.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    for row in COMPLIANCE_TABLE:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value

    document.save(path)


def make_questionnaire(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Security"
    sheet.append(["ID", "Domain", "Question", "Vendor response"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in QUESTIONNAIRE:
        sheet.append([*row, ""])
    sheet.column_dimensions["C"].width = 70
    sheet.column_dimensions["D"].width = 50
    for row in sheet.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True)
    workbook.save(path)


if __name__ == "__main__":
    import sys

    # The RFP and questionnaire are part of the frozen V0 corpus (baselines/v0/BASELINE.md). Office
    # files embed save timestamps, so rewriting them changes their bytes: never overwrite by default.
    force = "--force" in sys.argv
    builders = {
        "sample_rfp.docx": make_rfp,
        "sample_security_questionnaire.xlsx": make_questionnaire,
        **{name: (lambda path, t=title, s=sections: make_past_proposal(path, t, s))
           for name, (title, sections) in PAST_PROPOSALS.items()},
    }
    for filename, build in builders.items():
        path = HERE / filename
        if path.exists() and not force:
            print(f"kept    {filename} (exists; use --force to rewrite)")
            continue
        build(path)
        print(f"wrote   {filename}")
