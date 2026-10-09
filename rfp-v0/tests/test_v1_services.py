"""Application services without HTTP: library, retrieval, jobs, drafting, review, and export."""

from __future__ import annotations

import asyncio
import io

import docx
import openpyxl
import pytest
from sqlalchemy import select

from rfp_assistant.providers.base import LLMError
from rfp_assistant.schemas import DraftClaimOut, DraftResult
from rfp_assistant.errors import PipelineError
from rfp_assistant.api.v1 import library, projects
from rfp_assistant.api.v1.db import Answer, DraftRow, Job, Pair, PastProposal, Project, RequirementRow
from rfp_assistant.api.v1.export import FinalAnswer, export_docx, export_xlsx
from rfp_assistant.api.v1.retrieval import retrieve
from rfp_assistant.api.v1.sync import touch
from tests.conftest import docx_bytes
from tests.v1_fakes import FakeMemory, FakeV1LLM, make_context, pair, req

PAIRS = [
    pair("Which SSO standards do you support?", "We support SAML 2.0 and OpenID Connect with Entra ID and Okta.", reference="3.1"),
    pair("Describe automatic user provisioning.", "We support SCIM 2.0 provisioning with Okta and Entra ID.", reference="3.2"),
    pair("Provide pricing.", "See the commercial response.", reference="7.1"),
]


def run(coro):
    return asyncio.run(coro)


async def import_library(ctx, decisions=None):
    proposal, job = library.create_past_proposal(
        ctx, filename="past.docx", data=docx_bytes("3.1 Q", "Response: A"), client="Northwind Bank",
        industry="finance", submitted_on=None, result="won", loss_reason=None,
    )
    await ctx.jobs.wait(job.id)
    with ctx.db.session() as s:
        pairs = s.scalars(select(Pair).where(Pair.past_proposal_id == proposal.id).order_by(Pair.order)).all()
    decisions = decisions or [library.PairDecision(p.id, "dropped" if "pricing" in p.question.lower() else "kept") for p in pairs]
    created = library.confirm_pairs(ctx, proposal.id, decisions)
    await ctx.sync()
    return proposal, created


# --- library + outbox ---------------------------------------------------------------------------


def test_import_confirm_and_sync_to_memory(tmp_path):
    memory = FakeMemory()
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS), memory)

    async def main():
        proposal, created = await import_library(ctx)
        assert [a.code for a in created] == ["ANS-0001", "ANS-0002"]  # pricing pair dropped
        with ctx.db.session() as s:
            assert s.get(PastProposal, proposal.id).status == "confirmed"
            assert {a.hindsight_status for a in s.scalars(select(Answer))} == {"retained"}
        assert set(memory.docs) == {"ANS-0001", "ANS-0002"}
        assert memory.docs["ANS-0001"].client == "Northwind Bank"

    run(main())


def test_edited_pair_is_stored_as_edited(tmp_path):
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS[:1]))

    async def main():
        proposal, job = library.create_past_proposal(ctx, filename="p.docx", data=docx_bytes("x"), client=None,
                                                     industry=None, submitted_on=None, result="unknown", loss_reason=None)
        await ctx.jobs.wait(job.id)
        with ctx.db.session() as s:
            pair_id = s.scalars(select(Pair.id)).one()
        created = library.confirm_pairs(ctx, proposal.id, [library.PairDecision(pair_id, "edited", answer="Corrected answer.")])
        assert created[0].answer == "Corrected answer."

    run(main())


def test_confirm_requires_a_decision_for_every_pair_and_the_right_state(tmp_path):
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS))

    async def main():
        proposal, job = library.create_past_proposal(ctx, filename="p.docx", data=docx_bytes("x"), client=None,
                                                     industry=None, submitted_on=None, result="unknown", loss_reason=None)
        await ctx.jobs.wait(job.id)
        with pytest.raises(PipelineError) as error:
            library.confirm_pairs(ctx, proposal.id, [])
        assert error.value.http_status == 422
        with ctx.db.session() as s:
            ids = s.scalars(select(Pair.id)).all()
        library.confirm_pairs(ctx, proposal.id, [library.PairDecision(i, "kept") for i in ids])
        with pytest.raises(PipelineError) as error:
            library.confirm_pairs(ctx, proposal.id, [library.PairDecision(i, "kept") for i in ids])
        assert error.value.http_status == 409

    run(main())


def test_failed_pair_extraction_marks_proposal_and_job_failed(tmp_path):
    ctx = make_context(tmp_path, FakeV1LLM(fail_pairs=LLMError("refused", "pair extraction: the model declined")))

    async def main():
        proposal, job = library.create_past_proposal(ctx, filename="p.docx", data=docx_bytes("x"), client=None,
                                                     industry=None, submitted_on=None, result="unknown", loss_reason=None)
        await ctx.jobs.wait(job.id)
        with ctx.db.session() as s:
            assert s.get(PastProposal, proposal.id).status == "failed"
            job_row = s.get(Job, job.id)
            assert (job_row.status, job_row.error) == ("failed", "pair extraction: the model declined")

    run(main())


def test_memory_outage_keeps_answers_pending_until_a_later_sync(tmp_path):
    memory = FakeMemory()
    memory.available = False
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS[:2]), memory)

    async def main():
        await import_library(ctx)
        with ctx.db.session() as s:
            rows = s.scalars(select(Answer)).all()
            assert {r.hindsight_status for r in rows} == {"pending"}
            assert rows[0].hindsight_attempts == 1 and "down" in rows[0].hindsight_error
        memory.available = True
        report = await ctx.sync()
        assert (report.retained, report.pending) == (2, 0)

    run(main())


def test_deleted_answer_never_reaches_retrieval_even_if_hindsight_still_has_it(tmp_path):
    memory = FakeMemory()
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS[:2]), memory)

    async def main():
        await import_library(ctx)
        memory.available = False  # Hindsight can't be told yet
        library.delete_answer(ctx, "ANS-0002")
        memory.available = True
        assert "ANS-0002" in memory.docs  # stale copy still there...
        found = await retrieve("automatic user provisioning SCIM", memory=memory, db=ctx.db, k=3)
        assert "ANS-0002" not in [p.id for p in found.past_answers]  # ...but SQLite filters it out
        await ctx.sync()
        assert "ANS-0002" not in memory.docs

    run(main())


def test_an_edit_made_during_a_sync_is_not_lost(tmp_path):
    memory = FakeMemory()
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS[:1]), memory)

    async def main():
        await import_library(ctx)
        with ctx.db.session() as s:
            answer = s.scalars(select(Answer)).one()
            answer.answer, answer.hindsight_status = "Version 2", "pending"
            touch(answer)
            s.commit()

        def edit_during_retain(item):  # a reviewer saves version 3 while version 2 is being pushed
            with ctx.db.session() as s2:
                row = s2.scalars(select(Answer)).one()
                row.answer = "Version 3"
                touch(row)
                s2.commit()

        memory.on_retain = edit_during_retain
        await ctx.sync()
        memory.on_retain = None
        with ctx.db.session() as s:
            assert s.scalars(select(Answer)).one().hindsight_status == "pending"  # not wrongly marked done
        await ctx.sync()
        assert memory.docs["ANS-0001"].answer == "Version 3"

    run(main())


def test_retrieval_keeps_hindsight_order_limits_to_k_and_reports_outage(tmp_path):
    memory = FakeMemory()
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS[:2]), memory)

    async def main():
        await import_library(ctx)
        found = await retrieve("SCIM automatic user provisioning", memory=memory, db=ctx.db, k=1)
        assert [p.id for p in found.past_answers] == ["ANS-0002"]
        assert found.retrieved[0]["id"] == "ANS-0002" and found.retrieved[0]["position"] == 1
        memory.available = False
        down = await retrieve("anything", memory=memory, db=ctx.db, k=3)
        assert down.past_answers == [] and "unavailable" in down.warning

    run(main())


# --- projects: extraction, drafting, resume, review ---------------------------------------------

RFP_REQS = [
    req("Does your solution support SSO standards like SAML?", section="Security", reference="3.1"),
    req("Does your solution support automatic user provisioning via SCIM?", section="Security", reference="3.2"),
    req("Provide customer references.", section="Company", reference="2.2"),
]


async def project_ready(ctx):
    await import_library(ctx)
    project, job = projects.create_project(ctx, filename="rfp.docx", data=docx_bytes("x"), name="Harborview",
                                           client="Harborview CU", industry="finance")
    await ctx.jobs.wait(job.id)
    return project


def test_drafting_uses_library_answers_and_rejects_citations_it_was_not_shown(tmp_path):
    library_codes = ("ANS-0001", "ANS-0002")

    def drafter(r, past, instructions):
        offered = {p.id for p in past}
        if "references" in r.question:
            # Cite a real, live library answer that was NOT offered for this requirement.
            not_offered = next(c for c in library_codes if c not in offered)
            return DraftResult(answer="We have references.", claims=[DraftClaimOut(text="refs", source_ids=[not_offered])],
                               unsupported_claims=[], needs_sme=False, sme_question=None)
        top = past[0]
        return DraftResult(answer=top.answer, claims=[DraftClaimOut(text=top.answer, source_ids=[top.id, "FACT-003"])],
                           unsupported_claims=[], needs_sme=False, sme_question=None)

    llm = FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS, drafter=drafter)
    ctx = make_context(tmp_path, llm, retrieval_top_k=1)  # one past answer per requirement

    async def main():
        project = await project_ready(ctx)
        job = projects.start_drafting(ctx, project.id)
        await ctx.jobs.wait(job.id)
        with ctx.db.session() as s:
            p = s.get(Project, project.id)
            assert p.state == "in_review"
            drafts = {r.reference: r.drafts[-1] for r in p.requirements}
        assert drafts["3.2"].sources == ["ANS-0002", "FACT-003"] and drafts["3.2"].flags == []
        assert drafts["3.2"].retrieved[0]["id"] == "ANS-0002"
        assert all(len(offered) == 1 for _, offered, _ in llm.drafted)
        # The cited answer exists in the library, but wasn't shown for this requirement: invalid.
        assert "invalid_citation" in drafts["2.2"].flags
        assert drafts["2.2"].sources == []

    run(main())


def test_draft_all_resumes_after_a_crash_without_duplicates(tmp_path):
    llm = FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS)
    ctx = make_context(tmp_path, llm)

    async def main():
        project = await project_ready(ctx)
        with ctx.db.session() as s:
            p = s.get(Project, project.id)
            p.state = "drafting"
            first = p.requirements[0]
            job = Job(kind="draft_all", target_id=project.id, status="running", payload={})  # the process died here
            s.add(job)
            s.flush()
            s.add(DraftRow(requirement_id=first.id, job_id=job.id, version=1, status="drafted", answer="done before crash",
                           prompt_version="v1.0"))
            s.commit()
            job_id = job.id
        llm.drafted.clear()
        resumed = await ctx.jobs.resume_interrupted()
        assert job_id in resumed
        await ctx.jobs.wait(job_id)
        assert [d[0] for d in llm.drafted] == ["REQ-002", "REQ-003"]  # REQ-001 not redrafted
        with ctx.db.session() as s:
            job = s.get(Job, job_id)
            assert (job.status, job.done, job.total) == ("completed", 3, 3)
            assert len(s.scalars(select(DraftRow).where(DraftRow.job_id == job_id)).all()) == 3

    run(main())


def test_hindsight_outage_during_drafting_degrades_to_fact_sheet_only(tmp_path):
    memory = FakeMemory()
    llm = FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS)
    ctx = make_context(tmp_path, llm, memory)

    async def main():
        project = await project_ready(ctx)
        memory.available = False
        job = projects.start_drafting(ctx, project.id)
        await ctx.jobs.wait(job.id)
        with ctx.db.session() as s:
            assert "unavailable" in s.get(Job, job.id).warning
            assert all(r.drafts[-1].status == "needs_sme" for r in s.get(Project, project.id).requirements)
        assert all(d[1] == [] for d in llm.drafted)  # drafted with an empty library, not crashed

    run(main())


def test_drafting_without_company_facts_is_refused_and_leaves_the_project_as_it_was(tmp_path):
    llm = FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS)
    ctx = make_context(tmp_path, llm, fact_sheet_path=tmp_path / "no_such_facts" / "fact_sheet.json")

    async def main():
        project = await project_ready(ctx)
        with ctx.db.session() as s:
            state_before, jobs_before = s.get(Project, project.id).state, len(s.scalars(select(Job)).all())
        with pytest.raises(PipelineError) as raised:
            projects.start_drafting(ctx, project.id)
        assert (raised.value.code, raised.value.http_status) == ("no_company_facts", 409)
        with ctx.db.session() as s:
            assert s.get(Project, project.id).state == state_before != "drafting"  # not stuck, so it can be retried
            assert len(s.scalars(select(Job)).all()) == jobs_before  # no job was created
        assert llm.drafted == []

    run(main())


def test_review_actions_update_the_library_and_project_state(tmp_path):
    memory = FakeMemory()
    llm = FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS[:2])
    ctx = make_context(tmp_path, llm, memory)

    async def main():
        project = await project_ready(ctx)
        job = projects.start_drafting(ctx, project.id)
        await ctx.jobs.wait(job.id)
        with ctx.db.session() as s:
            r1, r2 = s.get(Project, project.id).requirements

        projects.review(ctx, r1.id, "accepted", None)
        with pytest.raises(PipelineError):
            projects.review(ctx, r2.id, "edited", "   ")  # an edit needs text
        projects.review(ctx, r2.id, "edited", "Our edited SCIM answer.")
        await ctx.sync()
        with ctx.db.session() as s:
            assert s.get(Project, project.id).state == "approved"
            project_answers = s.scalars(select(Answer).where(Answer.source == "project")).all()
            assert sorted(a.answer for a in project_answers)[0] == "Our edited SCIM answer."
            code = [a.code for a in project_answers if a.requirement_id == r2.id][0]
        assert memory.docs[code].answer == "Our edited SCIM answer."

        projects.review(ctx, r2.id, "rewritten", "A complete rewrite.")  # updates the same library answer
        projects.review(ctx, r1.id, "rejected", None)  # removes r1's library answer and reopens the project
        await ctx.sync()
        with ctx.db.session() as s:
            assert s.get(Project, project.id).state == "in_review"
            rows = s.scalars(select(Answer).where(Answer.source == "project")).all()
            assert {(a.requirement_id, a.status) for a in rows} == {(r1.id, "deleted"), (r2.id, "approved")}
            assert [a.answer for a in rows if a.status == "approved"] == ["A complete rewrite."]

    run(main())


def test_accepting_an_empty_draft_is_refused(tmp_path):
    llm = FakeV1LLM(pairs=[], requirements=RFP_REQS[2:])
    ctx = make_context(tmp_path, llm)

    async def main():
        project, job = projects.create_project(ctx, filename="rfp.docx", data=docx_bytes("x"), name="N", client=None, industry=None)
        await ctx.jobs.wait(job.id)
        job = projects.start_drafting(ctx, project.id)
        await ctx.jobs.wait(job.id)
        with ctx.db.session() as s:
            only = s.get(Project, project.id).requirements[0]
            assert only.drafts[-1].status == "needs_sme"
            assert "SME COMPLETION TEMPLATE" in only.drafts[-1].answer
            assert "sme_template" in only.drafts[-1].flags
        with pytest.raises(PipelineError) as error:
            projects.review(ctx, only.id, "accepted", None)
        assert error.value.http_status == 422
        with pytest.raises(PipelineError):
            projects.review(ctx, only.id, "edited", only.drafts[-1].answer)

    run(main())


def test_regenerate_creates_a_new_version_with_instructions_and_reopens_review(tmp_path):
    llm = FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS[:1])
    ctx = make_context(tmp_path, llm)

    async def main():
        project = await project_ready(ctx)
        await ctx.jobs.wait(projects.start_drafting(ctx, project.id).id)
        with ctx.db.session() as s:
            r1 = s.get(Project, project.id).requirements[0]
        projects.review(ctx, r1.id, "accepted", None)
        draft = await projects.regenerate(ctx, r1.id, "Make it shorter")
        assert draft.version == 2 and draft.instructions == "Make it shorter"
        assert llm.drafted[-1][2] == "Make it shorter"
        with ctx.db.session() as s:
            row = s.get(RequirementRow, r1.id)
            assert not projects.is_final(row)
            assert s.get(Project, project.id).state == "in_review"

    run(main())


def test_requirements_can_only_be_replaced_before_drafting(tmp_path):
    ctx = make_context(tmp_path, FakeV1LLM(pairs=PAIRS[:2], requirements=RFP_REQS))

    async def main():
        project = await project_ready(ctx)
        projects.replace_requirements(ctx, project.id, [projects.RequirementEdit(question="Only one question now?")])
        with ctx.db.session() as s:
            assert [r.question for r in s.get(Project, project.id).requirements] == ["Only one question now?"]
        await ctx.jobs.wait(projects.start_drafting(ctx, project.id).id)
        with pytest.raises(PipelineError) as error:
            projects.replace_requirements(ctx, project.id, [projects.RequirementEdit(question="Too late")])
        assert error.value.http_status == 409

    run(main())


# --- export -------------------------------------------------------------------------------------

ANSWERS = [
    FinalAnswer(code="REQ-001", section="Security", reference="SEC-01", question="Do you support SAML 2.0 single sign-on?", answer="Yes, SAML 2.0."),
    FinalAnswer(code="REQ-002", section="Security", reference="SEC-99", question="A question not in the workbook?", answer="Answer B."),
]


def test_docx_export_contains_every_answer_in_order():
    content = export_docx("Harborview: response", "Response prepared for Harborview CU", ANSWERS)
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(content)).paragraphs)
    assert text.index("Yes, SAML 2.0.") < text.index("Answer B.")
    assert "SEC-01 Do you support SAML 2.0 single sign-on?" in text


def test_xlsx_export_fills_the_response_column_and_reports_unplaced(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Question", "Vendor response"])
    sheet.append(["SEC-01", "Do you support SAML 2.0 single sign-on?", None])
    path = tmp_path / "q.xlsx"
    workbook.save(path)

    content, unplaced = export_xlsx(path, ANSWERS)
    filled = openpyxl.load_workbook(io.BytesIO(content))
    assert filled.active.cell(row=2, column=3).value == "Yes, SAML 2.0."
    assert unplaced == ["REQ-002"]
    assert filled["Unplaced answers"].cell(row=2, column=3).value == "Answer B."
