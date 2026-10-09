"""Build the upload-ready Virtusa demo pack.

Company facts come from the public Virtusa sources in corpus.json. Every client, proposal,
outcome and engagement detail is explicitly synthetic.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
CORPUS = HERE / "corpus.json"
NL = chr(10)


def _prepare(path: Path, force: bool) -> None:
    if path.exists() and not force:
        raise SystemExit(f"{path.name} already exists; rerun with --force to replace generated files.")


def _base_document(title: str, disclaimer: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    heading = document.add_heading(title, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning = document.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = warning.add_run(disclaimer)
    run.bold = True
    run.font.color.rgb = RGBColor(180, 35, 45)
    return document


def _metadata(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[0].paragraphs[0].runs[0].bold = True
    document.add_paragraph()


def _proposal_doc(proposal: dict, force: bool) -> None:
    path = HERE / proposal["filename"]
    _prepare(path, force)
    document = _base_document(proposal["title"], proposal["disclaimer"])
    _metadata(document, [
        ("Responding organization", "Virtusa"),
        ("Prospective client", proposal["client"]),
        ("Industry", proposal["industry"]),
        ("Submission date", proposal["submitted_on"]),
        ("Dataset status", "Synthetic demonstration history"),
    ])
    for section in proposal["sections"]:
        document.add_heading(section["name"], level=1)
        for pair in section["pairs"]:
            q = document.add_paragraph()
            q.add_run(f"{pair['reference']} Question: ").bold = True
            q.add_run(pair["question"])
            a = document.add_paragraph()
            a.add_run("Response: ").bold = True
            a.add_run(pair["answer"])
    footer = document.add_paragraph("End of synthetic demonstration proposal.")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(path)


def _rfp_doc(rfp: dict, force: bool) -> None:
    path = HERE / rfp["filename"]
    _prepare(path, force)
    document = _base_document(rfp["title"], rfp["disclaimer"])
    _metadata(document, [
        ("Issuing organization", rfp["client"]),
        ("Industry", rfp["industry"]),
        ("Document type", "Synthetic demonstration RFP"),
    ])
    document.add_heading("Response instructions", level=1)
    document.add_paragraph(rfp["instructions"])
    current = None
    for item in rfp["requirements"]:
        if item["section"] != current:
            current = item["section"]
            document.add_heading(current, level=1)
        p = document.add_paragraph()
        p.add_run(f"{item['reference']} ").bold = True
        p.add_run(item["question"])
        constraints = []
        if item.get("mandatory"):
            constraints.append("Mandatory")
        if item.get("word_limit"):
            constraints.append(f"Maximum {item['word_limit']} words")
        if constraints:
            note = document.add_paragraph(" | ".join(constraints))
            note.style = document.styles["Caption"]
    document.save(path)


def _write_support_files(data: dict, force: bool) -> None:
    facts_path = HERE / "virtusa_fact_sheet.json"
    text_path = HERE / "virtusa_company_facts.txt"
    sources_path = HERE / "SOURCES.md"
    manifest_path = HERE / "upload_manifest.json"
    readme_path = HERE / "README.md"
    for path in (facts_path, text_path, sources_path, manifest_path, readme_path):
        _prepare(path, force)

    clean_facts = [{k: v for k, v in fact.items() if k != "source"} for fact in data["facts"]]
    facts_path.write_text(json.dumps({"company": data["company"], "facts": clean_facts}, indent=2), encoding="utf-8")
    text_path.write_text(
        NL.join(f"{fact['topic']}: {fact['statement']}" for fact in data["facts"]) + NL,
        encoding="utf-8",
    )

    source_lines = [
        "# Virtusa fact source register",
        "",
        f"Reviewed on {data['reviewed_on']}. {data['disclaimer']}",
        "",
        "| Fact | Topic | Official source | Review by |",
        "|---|---|---|---|",
    ]
    for fact in data["facts"]:
        source = data["sources"][fact["source"]]
        review_by = fact["valid_to"] or "No automatic expiry; still review before a real bid"
        source_lines.append(f"| {fact['id']} | {fact['topic']} | [{source['title']}]({source['url']}) | {review_by} |")
    sources_path.write_text(NL.join(source_lines) + NL, encoding="utf-8")

    proposals = data["proposals"] + data["lost_proposals"]
    manifest = {
        "disclosure": data["disclaimer"],
        "company_fact_sheet": facts_path.name,
        "past_proposals": [
            {k: proposal[k] for k in ("filename", "client", "industry", "submitted_on", "recommended_result", "loss_reason")}
            for proposal in proposals
        ],
        "target_project": {
            "filename": data["target_rfp"]["filename"],
            "name": data["target_rfp"]["title"],
            "client": data["target_rfp"]["client"],
            "industry": data["target_rfp"]["industry"],
        },
    }
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    readme = """# Virtusa demo upload pack

**Disclosure:** Public Virtusa facts are separated from synthetic proposal history. The fictional
clients, outcomes and delivery examples must never be represented as real Virtusa engagements.

## Upload order

1. Import virtusa_fact_sheet.json on Company and save the facts.
2. Upload the four historical proposal DOCX files in the order shown in upload_manifest.json; enter
   the listed metadata and outcomes.
3. Confirm the extracted question/answer pairs and let both Hindsight banks finish syncing.
4. Upload the target RFP DOCX as a new project. Correct its requirements before running a comparison.
5. Run Does memory change the draft?, then draft, review, export, record an outcome and add a debrief.

The two won proposals are intentionally specific. The two lost proposals are intentionally vague
but not false. The target RFP is worded like the weak proposals so plain semantic retrieval has a
fair chance to choose them, while outcome memory can prefer equally relevant successful answers.
"""
    readme_path.write_text(readme, encoding="utf-8")


def validate(data: dict) -> None:
    facts = data["facts"]
    assert len(facts) == 18
    assert len({fact["id"] for fact in facts}) == len(facts)
    assert all(fact["source"] in data["sources"] for fact in facts)
    proposals = data["proposals"] + data["lost_proposals"]
    assert len(proposals) == 4
    assert sum(p["recommended_result"] == "won" for p in proposals) == 2
    assert sum(p["recommended_result"] == "lost" for p in proposals) == 2
    assert all("SYNTHETIC DEMO" in p["disclaimer"] for p in proposals)
    assert all(sum(len(s["pairs"]) for s in p["sections"]) == 6 for p in proposals)
    assert len(data["target_rfp"]["requirements"]) == 6


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    data = json.loads(CORPUS.read_text(encoding="utf-8"))
    validate(data)
    _write_support_files(data, args.force)
    for proposal in data["proposals"] + data["lost_proposals"]:
        _proposal_doc(proposal, args.force)
    _rfp_doc(data["target_rfp"], args.force)
    print("Built Virtusa demo pack: 18 facts, 4 historical proposals, 1 target RFP.")


if __name__ == "__main__":
    main()
